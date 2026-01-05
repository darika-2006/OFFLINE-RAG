from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from typing import List
from sentence_transformers import SentenceTransformer
import pypdfium2 as pdfium
from docx import Document
import psycopg2
import faiss
import numpy as np
import os
import json
import pickle

# ---------------- APP ----------------
app = FastAPI()

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

UPLOAD_DIR = "uploads"
INDEX_DIR = "index"
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(INDEX_DIR, exist_ok=True)

# ---------------- DB ----------------
conn = psycopg2.connect(
    dbname="RAG",
    user="postgres",
    password="darika_2006",
    host="localhost",
    port="5432"
)
cur = conn.cursor()

# Validate connection
try:
    cur.execute("SELECT 1")
    print("✅ Database connection successful")
except Exception as e:
    print(f"❌ Database connection failed: {e}")
    raise

# ---------------- MODEL ----------------
model = SentenceTransformer("sentence-transformers/all-mpnet-base-v2")

# ---------------- LOAD OR CREATE FAISS INDEX ----------------
FAISS_INDEX_PATH = os.path.join(INDEX_DIR, "hal_index.faiss")
MAPPING_PATH = os.path.join(INDEX_DIR, "chunk_id_mapping.pkl")

# Load or create FAISS index
if os.path.exists(FAISS_INDEX_PATH):
    index = faiss.read_index(FAISS_INDEX_PATH)
    print(f"✅ Loaded FAISS index with {index.ntotal} vectors")
else:
    index = faiss.IndexFlatL2(768)
    print("✅ Created new empty FAISS index")

# Load or create mapping: faiss_position → chunk_id
if os.path.exists(MAPPING_PATH):
    with open(MAPPING_PATH, 'rb') as f:
        chunk_id_mapping = pickle.load(f)
    print(f"✅ Loaded chunk_id mapping with {len(chunk_id_mapping)} entries")
else:
    chunk_id_mapping = []  # List where index = FAISS position, value = chunk_id
    print("✅ Created new empty mapping")

# Verify alignment
if len(chunk_id_mapping) != index.ntotal:
    print(f"⚠️ WARNING: Mapping size ({len(chunk_id_mapping)}) != FAISS size ({index.ntotal})")
    print("   Rebuilding mapping from database...")
    
    # Rebuild mapping from database
    cur.execute("SELECT chunk_id FROM chunks_metadata WHERE is_active = TRUE ORDER BY chunk_id ASC")
    chunk_id_mapping = [row[0] for row in cur.fetchall()]
    
    if len(chunk_id_mapping) != index.ntotal:
        print(f"❌ ERROR: DB has {len(chunk_id_mapping)} chunks but FAISS has {index.ntotal} vectors")
        print("   Run rebuild_faiss.py to fix alignment")

# ---------------- ROOT ENDPOINT ----------------
@app.get("/")
def root():
    """Health check endpoint"""
    return {
        "status": "online",
        "service": "Upload & Embedding Service",
        "faiss_index_exists": os.path.exists(FAISS_INDEX_PATH),
        "faiss_index_size": index.ntotal,
        "mapping_size": len(chunk_id_mapping),
        "alignment_ok": len(chunk_id_mapping) == index.ntotal
    }

# ---------------- TEXT EXTRACTION ----------------
def extract_text_from_pdf(pdf_path, max_words=350, overlap=70):
    pdf = pdfium.PdfDocument(pdf_path)
    chunks = []

    for page_num in range(len(pdf)):
        page = pdf[page_num]
        text = page.get_textpage().get_text_range()
        
        words = text.split()
        start = 0
        
        while start < len(words):
            end = start + max_words
            chunk_text = " ".join(words[start:end])
            
            chunks.append({
                "text": chunk_text,
                "page_number": page_num + 1
            })
            
            start += (max_words - overlap)

    pdf.close()
    return chunks


def extract_text_from_docx(docx_path, max_words=250, overlap=50):
    doc = Document(docx_path)
    full_text = " ".join([p.text for p in doc.paragraphs])

    words = full_text.split()
    chunks = []
    start = 0
    page_number = 1

    while start < len(words):
        end = start + max_words
        chunk_text = " ".join(words[start:end])
        
        chunks.append({
            "text": chunk_text,
            "page_number": page_number
        })
        
        start += (max_words - overlap)
        if start >= len(words):
            break
        page_number += 1

    return chunks


# ---------------- CHUNK + DB ----------------
def chunk_list(file_paths, allowed_roles, uploaded_by):
    chunks = []

    for path in file_paths:
        # Check if uploaded_by is user_id or username
        cur.execute("SELECT user_id FROM login_credentials WHERE user_id = %s", (uploaded_by,))
        result = cur.fetchone()
        
        if result:
            user_id = result[0]
            print(f"✅ Found user by user_id: {user_id}")
        else:
            cur.execute("SELECT user_id FROM login_credentials WHERE username = %s", (uploaded_by,))
            result = cur.fetchone()
            
            if result:
                user_id = result[0]
                print(f"✅ Found user by username: {uploaded_by} -> user_id: {user_id}")
            else:
                cur.execute("SELECT user_id, username FROM login_credentials")
                available_users = cur.fetchall()
                user_list = ", ".join([f"{u[0]} ({u[1]})" for u in available_users])
                
                raise HTTPException(
                    status_code=400,
                    detail=f"User '{uploaded_by}' not found. Available: {user_list}"
                )
        
        # Insert document
        doc_name = os.path.basename(path)
        sql_doc = """
        INSERT INTO documents(doc_name, doc_path, allowed_roles, uploaded_by)
        VALUES (%s, %s, %s, %s) RETURNING doc_id
        """
        cur.execute(sql_doc, (doc_name, path, allowed_roles, user_id))
        doc_id = cur.fetchone()[0]
        conn.commit()
        
        print(f"✅ Inserted document: {doc_name} (doc_id: {doc_id})")

        # Extract text
        if path.lower().endswith('.pdf'):
            temp_chunks = extract_text_from_pdf(path)
        elif path.lower().endswith('.docx'):
            temp_chunks = extract_text_from_docx(path)
        else:
            print(f"⚠️ Skipping unsupported file: {path}")
            continue

        # Add metadata
        for chunk in temp_chunks:
            chunk["doc_id"] = doc_id
            chunk["allowed_roles"] = allowed_roles
            chunks.append(chunk)

    return chunks


def build_metadata(chunks):
    """
    Insert chunks and return list of chunk_ids in insertion order
    """
    chunk_ids = []
    
    for chunk in chunks:
        sql_chunk = """
        INSERT INTO chunks_metadata(doc_id, page_number, chunk_text, allowed_roles)
        VALUES (%s, %s, %s, %s)
        RETURNING chunk_id
        """
        cur.execute(sql_chunk, (
            chunk["doc_id"],
            chunk["page_number"],
            chunk["text"],
            chunk["allowed_roles"]
        ))
        chunk_id = cur.fetchone()[0]
        chunk_ids.append(chunk_id)
    
    conn.commit()
    print(f"✅ Inserted {len(chunks)} chunks into database")
    
    return chunk_ids


# ---------------- UPLOAD API ----------------
@app.post("/upload")
async def upload_documents(
    files: List[UploadFile] = File(...),
    allowed_roles: str = Form(...),
    uploaded_by: str = Form(...)
):
    """
    Upload PDF/DOCX files to RAG system
    """
    
    try:
        if not files:
            raise HTTPException(status_code=400, detail="No files uploaded")
        
        print(f"\n📤 Upload request received:")
        print(f"   Files: {[f.filename for f in files]}")
        print(f"   Roles: {allowed_roles}")
        print(f"   User: {uploaded_by}")
        
        saved_paths = []
        
        # Save uploaded files
        for file in files:
            if not file.filename.endswith((".pdf", ".docx")):
                raise HTTPException(
                    status_code=400,
                    detail=f"Only PDF and DOCX files allowed. Got: {file.filename}"
                )
            
            file_path = os.path.join(UPLOAD_DIR, file.filename)
            
            with open(file_path, "wb") as f:
                content = await file.read()
                f.write(content)
            
            saved_paths.append(file_path)
            print(f"   ✅ Saved: {file_path}")
        
        # Process files
        print(f"\n🔄 Processing files...")
        chunks = chunk_list(saved_paths, allowed_roles.lower(), uploaded_by)
        
        if not chunks:
            raise HTTPException(status_code=400, detail="No chunks extracted from files")
        
        print(f"   ✅ Extracted {len(chunks)} chunks")
        
        # Store metadata in PostgreSQL and get chunk_ids
        new_chunk_ids = build_metadata(chunks)
        
        # Create embeddings
        print(f"\n🧮 Creating embeddings...")
        texts = [chunk["text"] for chunk in chunks]
        embeddings = model.encode(texts, normalize_embeddings=True)
        print(f"   ✅ Created {len(embeddings)} embeddings")
        
        # Update FAISS index
        index.add(embeddings.astype('float32'))
        
        # Update mapping: append new chunk_ids
        chunk_id_mapping.extend(new_chunk_ids)
        
        # Save both index and mapping
        faiss.write_index(index, FAISS_INDEX_PATH)
        with open(MAPPING_PATH, 'wb') as f:
            pickle.dump(chunk_id_mapping, f)
        
        print(f"   ✅ Updated FAISS index (now {index.ntotal} vectors)")
        print(f"   ✅ Updated mapping (now {len(chunk_id_mapping)} entries)")
        
        # Verify alignment
        if len(chunk_id_mapping) != index.ntotal:
            print(f"   ⚠️ WARNING: Mapping ({len(chunk_id_mapping)}) != FAISS ({index.ntotal})")
        
        response_data = {
            "message": "Upload successful",
            "files_uploaded": len(files),
            "total_chunks": len(chunks),
            "faiss_index_size": index.ntotal,
            "mapping_size": len(chunk_id_mapping)
        }
        
        print(f"\n✅ Upload completed successfully!")
        print(f"   Response: {response_data}")
        
        return response_data
    
    except HTTPException:
        raise
    except Exception as e:
        print(f"\n❌ Upload error: {str(e)}")
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Upload failed: {str(e)}")


# ---------------- DELETE API ----------------
@app.post("/delete-document")
def delete_document(doc_id: int):
    """
    Soft-delete a document - marks as inactive but keeps FAISS intact
    """
    try:
        sql_doc = "UPDATE documents SET is_active = FALSE WHERE doc_id = %s"
        sql_chunks = "UPDATE chunks_metadata SET is_active = FALSE WHERE doc_id = %s"
        
        cur.execute(sql_doc, (doc_id,))
        cur.execute(sql_chunks, (doc_id,))
        conn.commit()
        
        print(f"✅ Soft-deleted document {doc_id}")
        print(f"   Note: FAISS index not rebuilt. Deleted chunks will be filtered in queries.")
        
        return {"message": f"Document {doc_id} soft-deleted. FAISS index unchanged."}
    
    except Exception as e:
        print(f"❌ Delete error: {e}")
        raise HTTPException(status_code=500, detail=f"Delete failed: {str(e)}")
